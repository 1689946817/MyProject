from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
ASSET_DIR = DOCS_DIR / "midterm_ppt_assets" / "docx_media"
DOCX_PATH = DOCS_DIR / "张琪-中期报告-v2.docx"
TEMPLATE_PATH = DOCS_DIR / "PPT模版.pptx"
OUTPUT_PATH = DOCS_DIR / "张琪-中期汇报-docx版.pptx"

COLOR_NAVY = RGBColor(18, 55, 109)
COLOR_RED = RGBColor(183, 47, 47)
COLOR_TEXT = RGBColor(40, 40, 40)
COLOR_MUTED = RGBColor(100, 100, 100)
COLOR_LIGHT = RGBColor(244, 247, 252)
COLOR_BORDER = RGBColor(216, 223, 235)
COLOR_WHITE = RGBColor(255, 255, 255)


def ensure_docx_media() -> list[Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    with ZipFile(DOCX_PATH) as archive:
        media_names = sorted(
            [name for name in archive.namelist() if name.startswith("word/media/")],
            key=lambda item: int(Path(item).stem.replace("image", "")),
        )
        paths: list[Path] = []
        for name in media_names:
            target = ASSET_DIR / Path(name).name
            if not target.exists():
                target.write_bytes(archive.read(name))
            paths.append(target)
    return paths


def remove_all_slides(prs: Presentation) -> None:
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def add_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[0])


def add_top_banner(slide, title: str, subtitle: str | None = None) -> None:
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        Inches(10),
        Inches(0.7),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_NAVY
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.14), Inches(7.9), Inches(0.35))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COLOR_WHITE

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(8.0), Inches(0.18), Inches(1.7), Inches(0.28))
        tf = sub_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = subtitle
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(227, 232, 243)


def add_footer(slide, text: str) -> None:
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.45),
        Inches(7.15),
        Inches(9.1),
        Inches(0.02),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_RED
    line.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.2), Inches(5.5), Inches(0.18))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_MUTED


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    font_size: int = 18,
    bold: bool = False,
    color: RGBColor = COLOR_TEXT,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bullets(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    items: list[str],
    font_size: int = 18,
    color: RGBColor = COLOR_TEXT,
    spacing: int = 10,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.bullet = True


def add_card(slide, left: float, top: float, width: float, height: float, title: str, bullets: list[str]) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_LIGHT
    shape.line.color.rgb = COLOR_BORDER

    add_textbox(slide, left + 0.15, top + 0.12, width - 0.3, 0.28, title, font_size=16, bold=True, color=COLOR_NAVY)
    add_bullets(slide, left + 0.12, top + 0.42, width - 0.24, height - 0.5, bullets, font_size=13, spacing=4)


def contain_size(path: Path, max_w: float, max_h: float) -> tuple[float, float]:
    with Image.open(path) as img:
        width, height = img.size
    ratio = min(max_w / width, max_h / height)
    return width * ratio, height * ratio


def add_picture_contain(slide, path: Path, left: float, top: float, width: float, height: float) -> None:
    img_w, img_h = contain_size(path, Inches(width), Inches(height))
    x = Inches(left) + (Inches(width) - img_w) / 2
    y = Inches(top) + (Inches(height) - img_h) / 2
    slide.shapes.add_picture(str(path), x, y, width=img_w, height=img_h)


def add_image_frame(slide, left: float, top: float, width: float, height: float) -> None:
    frame = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = COLOR_WHITE
    frame.line.color.rgb = COLOR_BORDER


def add_caption(slide, left: float, top: float, width: float, text: str) -> None:
    add_textbox(slide, left, top, width, 0.24, text, font_size=10, color=COLOR_MUTED, align=PP_ALIGN.CENTER)


def build_title_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_textbox(slide, 0.75, 1.05, 7.8, 0.45, "中期汇报", font_size=20, bold=True, color=COLOR_RED)
    add_textbox(slide, 0.75, 1.65, 8.2, 1.1, "基于语义描述桥接的多模态RAG系统设计与实现", font_size=28, bold=True, color=COLOR_NAVY)
    add_textbox(slide, 0.8, 3.0, 7.8, 0.5, "以《张琪-中期报告-v2.docx》为唯一内容依据", font_size=16, color=COLOR_MUTED)
    add_textbox(slide, 0.85, 4.45, 4.5, 1.1, "汇报人：张琪\n学号：20225928\n班级：计算机2202", font_size=17)
    add_textbox(slide, 6.5, 5.05, 2.3, 0.6, "2026年4月", font_size=18, bold=True, color=COLOR_NAVY, align=PP_ALIGN.RIGHT)
    add_footer(slide, "东北大学毕业设计（论文）中期汇报")


def build_outline_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "汇报内容", "Outline")
    cards = [
        ("01 前期工作", ["文献阅读与资料学习", "系统设计与开发"]),
        ("02 系统展示", ["核心界面", "功能与优化点"]),
        ("03 实验评估", ["检索实验", "生成实验"]),
        ("04 微调工作", ["训练配置", "阶段性结果"]),
        ("05 问题分析", ["当前不足", "原因定位"]),
        ("06 后续计划", ["下一阶段安排", "论文材料准备"]),
    ]
    positions = [(0.7, 1.25), (3.55, 1.25), (6.4, 1.25), (0.7, 3.75), (3.55, 3.75), (6.4, 3.75)]
    for (title, bullets), (left, top) in zip(cards, positions):
        add_card(slide, left, top, 2.4, 1.95, title, bullets)
    add_footer(slide, "本版PPT按中期报告原文顺序重新组织，不引入外部实验图表")


def build_literature_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "前期文献阅读与资料学习", "Literature Review")
    add_bullets(
        slide,
        0.75,
        1.15,
        8.45,
        5.6,
        [
            "围绕多模态检索、检索增强生成、多模态大模型、向量数据库和 RAG 评测等方向进行了系统调研。",
            "重点学习了 CLIP、BLIP-2、LLaVA、Qwen-VL、MuRAG 以及多模态 RAG 综述与基准研究等代表性工作。",
            "同步查阅了 RAGAs 等自动评测方法，以及向量数据库、文本嵌入模型和检索系统实现相关资料。",
            "通过文献和技术资料学习，对检索相关性、答案忠实性、上下文利用和向量召回机制形成了初步认识。",
            "这些工作为后续系统设计、原型开发和实验评估提供了理论基础。",
        ],
        font_size=17,
    )
    add_footer(slide, "内容依据：中期报告“1.前期文献阅读与资料学习”")


def build_design_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统设计与开发", "Design & Development")
    add_card(slide, 0.7, 1.2, 2.7, 1.9, "总体设计", ["建立了业务模型、功能模型和数据模型。", "完成了功能设计、数据库设计和模块设计。"])
    add_card(slide, 3.65, 1.2, 2.7, 1.9, "后端技术栈", ["后端采用 FastAPI。", "围绕多模态 RAG 业务流程组织系统能力。"])
    add_card(slide, 6.6, 1.2, 2.7, 1.9, "前端技术栈", ["前端采用 Vue3 + TypeScript + Vite。", "形成可交互的演示型系统页面。"])
    add_bullets(
        slide,
        0.8,
        3.65,
        8.3,
        2.3,
        [
            "已完成图片知识库、文档知识库、图文检索服务、RAG 问答服务及基础管理功能开发。",
            "系统当前已经具备较完整的运行框架，可作为后续实验和论文撰写的核心原型。",
        ],
        font_size=17,
    )
    add_footer(slide, "内容依据：中期报告“2.系统设计与开发”")


def build_ui_top_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统界面展示（一）", "UI Screens")
    items = [
        ("图表 1 图片知识库界面", media[0]),
        ("图表 2 文档知识库界面", media[1]),
    ]
    positions = [(0.65, 1.05), (5.0, 1.05)]
    for (caption, path), (left, top) in zip(items, positions):
        add_image_frame(slide, left, top, 4.0, 4.85)
        add_picture_contain(slide, path, left + 0.08, top + 0.08, 3.84, 4.35)
        add_caption(slide, left, top + 4.45, 4.0, caption)
    add_footer(slide, "图片与标题均来自中期报告内嵌图表 1-2")


def build_ui_bottom_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统界面展示（二）", "UI Screens")
    items = [
        ("图表 3 图文检索服务界面", media[2]),
        ("图表 4 RAG 问答服务", media[3]),
    ]
    positions = [(0.65, 1.05), (5.0, 1.05)]
    for (caption, path), (left, top) in zip(items, positions):
        add_image_frame(slide, left, top, 4.0, 4.85)
        add_picture_contain(slide, path, left + 0.08, top + 0.08, 3.84, 4.35)
        add_caption(slide, left, top + 4.45, 4.0, caption)
    add_footer(slide, "图片与标题均来自中期报告内嵌图表 3-4")


def build_features_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "核心功能实现与优化", "Core Functions")
    add_card(slide, 0.7, 1.2, 2.7, 1.6, "图像语义化", ["实现图像描述生成、结构化信息入库和向量化存储。"])
    add_card(slide, 3.6, 1.2, 2.7, 1.6, "双路检索", ["实现文本检索图片和以图搜图两类检索机制。"])
    add_card(slide, 6.5, 1.2, 2.7, 1.6, "多模态问答", ["完成基于 RAG 的问答能力，可返回回答、来源和检索过程信息。"])
    add_card(slide, 0.7, 3.15, 2.7, 1.8, "检索增强", ["集成查询改写、多查询扩展、混合检索和重排序。"])
    add_card(slide, 3.6, 3.15, 2.7, 1.8, "上下文优化", ["完成上下文压缩与轻量化路由机制集成。"])
    add_card(slide, 6.5, 3.15, 2.7, 1.8, "可追溯性", ["能够根据问题进行意图识别并初步支持过程查看。"])
    add_footer(slide, "内容依据：中期报告“3.核心功能实现与优化”")


def build_eval_intro_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "实验评估总述", "Evaluation Summary")
    add_textbox(slide, 0.75, 1.0, 4.0, 0.35, "已完成两类实验评估", font_size=19, bold=True, color=COLOR_NAVY)
    add_bullets(
        slide,
        0.75,
        1.45,
        4.0,
        3.2,
        [
            "检索器评估",
            "生成器评估",
        ],
        font_size=18,
    )
    add_textbox(slide, 5.0, 1.0, 4.0, 0.35, "阶段性结论", font_size=19, bold=True, color=COLOR_NAVY)
    add_bullets(
        slide,
        5.0,
        1.45,
        4.0,
        3.4,
        [
            "proposed 方法在多个领域接近甚至超过 baseline_clip。",
            "在 CRM 和 energy 领域表现较好。",
            "在 finance 和 cross-domain 场景下仍有优化空间。",
            "生成实验初步验证了方法的有效性和竞争力。",
        ],
        font_size=16,
    )
    add_textbox(slide, 0.85, 5.1, 8.2, 1.0, "这一部分严格按中期报告原文陈述实验结果，图表全部来自报告内嵌图片，不额外引入其他实验目录中的版本。", font_size=16)
    add_footer(slide, "内容依据：中期报告“4.实验评估”")


def build_retrieval_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "检索实验结果", "Retrieval Evaluation")
    items = [
        ("图表 5 CRM领域相关检索指标", media[4]),
        ("图表 6 Energy领域相关检索指标", media[5]),
        ("图表 7 两种方法检索延迟对比", media[6]),
        ("图表 8 两种方法分领域指标热力图", media[7]),
    ]
    slots = [(0.45, 1.0), (5.0, 1.0), (0.45, 4.0), (5.0, 4.0)]
    for (caption, path), (left, top) in zip(items, slots):
        add_picture_contain(slide, path, left, top, 4.1, 2.35)
        add_caption(slide, left, top + 2.27, 4.1, caption)
    add_footer(slide, "内容依据：中期报告图表 5-8 与对应文字结论")


def build_generation_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "生成实验结果", "Generation Evaluation")
    items = [
        ("图表 9 三种方法生成器评测总览", media[8]),
        ("图表 10 三种方法分领域答案正确性", media[9]),
        ("图表 11 本文方法相对基线的提升", media[10]),
        ("图表 12 分问题类型答案正确性", media[11]),
        ("图表 13 相关性与忠实度对比", media[12]),
    ]
    placements = [
        (0.45, 1.0, 2.9, 2.15),
        (3.55, 1.0, 2.9, 2.15),
        (6.65, 1.0, 2.9, 2.15),
        (1.2, 4.0, 3.1, 2.15),
        (5.1, 4.0, 3.1, 2.15),
    ]
    for (caption, path), (left, top, width, height) in zip(items, placements):
        add_picture_contain(slide, path, left, top, width, height)
        add_caption(slide, left, top + height + 0.02, width, caption)
    add_footer(slide, "内容依据：中期报告图表 9-13 与对应文字结论")


def build_finetune_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "模型微调", "Fine-tuning")
    add_bullets(
        slide,
        0.7,
        1.15,
        4.0,
        4.8,
        [
            "已完成基于 UniDoc-Bench 图像及描述文本的监督微调数据构建、训练执行和阶段性评估。",
            "训练样本 888 条，验证样本 47 条。",
            "完成基于 Qwen3.5-4B 的 LoRA 微调训练，共 140 个 step，耗时约 5 小时 20 分钟。",
            "最终训练损失为 0.5194，验证损失为 0.5915；step-75 验证损失最低，为 0.5892。",
            "微调后在部分复杂页面上提升了格式遵循、关键词补充和局部细节描述能力。",
            "但仍存在输出冗长、重复等问题，后续还需继续优化数据质量和输出效果。",
        ],
        font_size=14,
        spacing=5,
    )
    add_picture_contain(slide, media[13], 5.0, 1.25, 4.0, 5.1)
    add_caption(slide, 5.05, 6.45, 3.9, "图表 14 微调学习曲线")
    add_footer(slide, "内容依据：中期报告“5.模型微调”")


def build_issues_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "存在问题", "Open Issues")
    add_card(slide, 0.75, 1.25, 4.0, 1.7, "1. OCR 基线检索异常", ["OCR 基线索引构建不完整，导致 Recall@K、MRR 等指标失真。"])
    add_card(slide, 5.0, 1.25, 4.0, 1.7, "2. 模型微调效果不够理想", ["当前微调虽完成阶段性训练与验证，但整体效果提升仍不够稳定。"])
    add_card(slide, 0.75, 3.25, 4.0, 1.7, "3. 系统响应速度仍有待提升", ["RAG 链路和多模态问答受串行检索优化策略和 IO 操作影响，延迟较大。"])
    add_card(slide, 5.0, 3.25, 4.0, 1.7, "4. 工程化与鲁棒性仍存在不足", ["异常处理、性能优化、恢复机制和工程化规范方面仍有提升空间。"])
    add_footer(slide, "内容依据：中期报告“存在问题”")


def build_plan_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "下一阶段工作计划", "Next Steps")
    steps = [
        ("01", "OCR 基线问题修复与重跑实验", "诊断并修复 OCR 索引构建问题，重新构建索引并重跑受影响评估。"),
        ("02", "继续优化模型微调与图像描述能力", "进一步探索 Qwen3.5 系列的 LoRA 参数与训练策略。"),
        ("03", "完善多模态 RAG 链路并提升系统性能", "继续优化检索、筛选、重排序和问答延迟。"),
        ("04", "加强测试并完善系统功能体验", "补充边界场景处理，优化前端交互和解析进度展示。"),
        ("05", "推进论文初稿撰写与材料整理", "同步整理系统架构图、流程图、界面截图和实验结果材料。"),
    ]
    top = 1.15
    for idx, (no, title, desc) in enumerate(steps):
        left = 0.7 if idx % 2 == 0 else 5.1
        row = idx // 2
        card_top = top + row * 1.85
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(card_top),
            Inches(3.8),
            Inches(1.45),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_LIGHT
        shape.line.color.rgb = COLOR_BORDER
        add_textbox(slide, left + 0.15, card_top + 0.12, 0.4, 0.25, no, font_size=18, bold=True, color=COLOR_RED)
        add_textbox(slide, left + 0.55, card_top + 0.12, 3.0, 0.25, title, font_size=15, bold=True, color=COLOR_NAVY)
        add_textbox(slide, left + 0.55, card_top + 0.46, 3.0, 0.6, desc, font_size=13)
    add_footer(slide, "内容依据：中期报告“下一阶段工作计划”")


def build_end_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_textbox(slide, 1.0, 2.15, 7.5, 0.8, "汇报完毕  谢谢老师", font_size=30, bold=True, color=COLOR_NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 2.2, 3.35, 5.2, 0.5, "张琪  |  东北大学计算机2202", font_size=18, color=COLOR_MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, 2.8, 4.05, 4.2, 0.4, "2026年4月", font_size=16, color=COLOR_RED, align=PP_ALIGN.CENTER)
    add_footer(slide, "中期汇报")


def main() -> None:
    media = ensure_docx_media()
    if len(media) != 14:
        raise ValueError(f"expected 14 docx media images, got {len(media)}")

    prs = Presentation(str(TEMPLATE_PATH))
    remove_all_slides(prs)

    build_title_slide(prs)
    build_outline_slide(prs)
    build_literature_slide(prs)
    build_design_slide(prs)
    build_ui_top_slide(prs, media)
    build_ui_bottom_slide(prs, media)
    build_features_slide(prs)
    build_eval_intro_slide(prs)
    build_retrieval_slide(prs, media)
    build_generation_slide(prs, media)
    build_finetune_slide(prs, media)
    build_issues_slide(prs)
    build_plan_slide(prs)
    build_end_slide(prs)

    prs.save(str(OUTPUT_PATH))
    print(f"saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
