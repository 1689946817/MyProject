from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
ASSET_DIR = DOCS_DIR / "midterm_ppt_assets" / "docx_media"
TEMPLATE_PATH = DOCS_DIR / "PPT模版.pptx"
OUTPUT_PATH = DOCS_DIR / "张琪-中期汇报.pptx"

RETRIEVAL_FIG_DIR = ROOT / "实验数据" / "检索器" / "analysis" / "figures_zh"
GEN_FIG_DIR = ROOT / "实验数据" / "生成器" / "analysis" / "figures_zh"
FINETUNE_FIG_DIR = ROOT / "实验数据" / "微调" / "analysis" / "figures_zh"

COLOR_NAVY = RGBColor(18, 55, 109)
COLOR_RED = RGBColor(183, 47, 47)
COLOR_TEXT = RGBColor(40, 40, 40)
COLOR_MUTED = RGBColor(100, 100, 100)
COLOR_LIGHT = RGBColor(244, 247, 252)
COLOR_BORDER = RGBColor(216, 223, 235)


def remove_all_slides(prs: Presentation) -> None:
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def add_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[0])


def add_top_banner(slide, title: str, subtitle: str | None = None) -> None:
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        Inches(10),
        Inches(0.7),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_NAVY
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.14), Inches(7.9), Inches(0.35))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(8.0), Inches(0.18), Inches(1.7), Inches(0.28))
        tf = sub_box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        run = p.add_run()
        run.text = subtitle
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(227, 232, 243)


def add_footer(slide, text: str) -> None:
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.45),
        Inches(7.15),
        Inches(9.1),
        Inches(0.02),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_RED
    line.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.2), Inches(5.5), Inches(0.18))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_MUTED


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    font_size: int = 18,
    bold: bool = False,
    color: RGBColor = COLOR_TEXT,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bullets(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    items: Iterable[str],
    font_size: int = 18,
    color: RGBColor = COLOR_TEXT,
    level: int = 0,
    spacing: int = 10,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = level
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.bullet = True


def add_card(slide, left: float, top: float, width: float, height: float, title: str, bullets: list[str]) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_LIGHT
    shape.line.color.rgb = COLOR_BORDER

    add_textbox(slide, left + 0.15, top + 0.12, width - 0.3, 0.28, title, font_size=16, bold=True, color=COLOR_NAVY)
    add_bullets(slide, left + 0.12, top + 0.42, width - 0.24, height - 0.5, bullets, font_size=13, spacing=4)


def contain_size(path: Path, max_w: float, max_h: float) -> tuple[float, float]:
    with Image.open(path) as img:
        width, height = img.size
    ratio = min(max_w / width, max_h / height)
    return width * ratio, height * ratio


def add_picture_contain(slide, path: Path, left: float, top: float, width: float, height: float) -> None:
    img_w, img_h = contain_size(path, Inches(width), Inches(height))
    x = Inches(left) + (Inches(width) - img_w) / 2
    y = Inches(top) + (Inches(height) - img_h) / 2
    slide.shapes.add_picture(str(path), x, y, width=img_w, height=img_h)


def add_caption(slide, left: float, top: float, width: float, text: str) -> None:
    add_textbox(slide, left, top, width, 0.22, text, font_size=10, color=COLOR_MUTED, align=PP_ALIGN.CENTER)


def build_title_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_textbox(slide, 0.75, 1.05, 7.8, 0.45, "中期汇报", font_size=20, bold=True, color=COLOR_RED)
    add_textbox(slide, 0.75, 1.65, 8.2, 1.1, "基于语义描述桥接的多模态RAG系统设计与实现", font_size=28, bold=True, color=COLOR_NAVY)
    add_textbox(slide, 0.8, 3.0, 7.8, 0.5, "Multimodal RAG System via Semantic Description Bridging", font_size=16, color=COLOR_MUTED)
    add_textbox(slide, 0.85, 4.45, 4.5, 1.1, "汇报人：张琪\n学号：20225928\n班级：计算机2202", font_size=17)
    add_textbox(slide, 6.5, 5.05, 2.3, 0.6, "2026年4月", font_size=18, bold=True, color=COLOR_NAVY, align=PP_ALIGN.RIGHT)
    add_footer(slide, "东北大学毕业设计（论文）中期汇报")


def build_outline_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "汇报内容", "Midterm Outline")
    cards = [
        ("01 课题背景", ["问题背景", "技术路线选择"]),
        ("02 阶段工作", ["文献调研", "系统设计与实现"]),
        ("03 系统现状", ["前后端功能闭环", "关键优化点"]),
        ("04 实验结果", ["检索实验", "生成实验"]),
        ("05 微调探索", ["LoRA 训练", "阶段性观察"]),
        ("06 问题与计划", ["当前不足", "下一阶段安排"]),
    ]
    positions = [(0.7, 1.25), (3.55, 1.25), (6.4, 1.25), (0.7, 3.75), (3.55, 3.75), (6.4, 3.75)]
    for (title, bullets), (left, top) in zip(cards, positions):
        add_card(slide, left, top, 2.4, 1.95, title, bullets)
    add_footer(slide, "整体结构以《张琪-中期报告-v2.docx》为纲，并结合当前代码与实验材料校准")


def build_background_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "课题背景与技术路线", "Background")
    add_textbox(slide, 0.7, 1.0, 4.4, 0.3, "选题背景", font_size=18, bold=True, color=COLOR_NAVY)
    add_bullets(
        slide,
        0.7,
        1.35,
        4.2,
        4.8,
        [
            "图像、图表、页面截图和多模态文档快速增长，传统文本RAG难以直接利用视觉语义。",
            "直接跨模态 embedding 方案链路短，但在复杂文档、抽象语义和可解释性方面存在局限。",
            "毕业设计目标是构建可运行、可验证、可展示的多模态RAG原型，为后续论文实验提供支撑。",
        ],
        font_size=16,
    )
    add_textbox(slide, 5.35, 1.0, 3.9, 0.3, "技术路线", font_size=18, bold=True, color=COLOR_NAVY)
    add_card(slide, 5.35, 1.42, 3.6, 1.15, "方案 A：统一跨模态向量空间", ["适合视觉相似检索", "强基线：baseline_clip"])
    add_card(slide, 5.35, 2.78, 3.6, 1.65, "方案 B：语义描述桥接（主方案）", ["MLLM 先生成结构化描述", "统一走文本 embedding + 检索 + RAG 链路", "更利于检索增强、来源回溯与论文论证"])
    add_textbox(slide, 5.35, 4.75, 3.6, 0.75, "当前主系统采用“图像/文档语义化 -> 文本检索 -> RAG问答”，同时保留 CLIP 类方法作为强基线对照。", font_size=15)
    add_footer(slide, "技术路线选择兼顾工程可落地性、可解释性和实验对比价值")


def build_progress_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "前一阶段工作总览", "Progress")
    add_card(slide, 0.65, 1.25, 2.2, 2.15, "文献与技术调研", ["梳理多模态检索、多模态RAG、评测方法和向量库方案", "形成主方案与基线方案的选型依据"])
    add_card(slide, 3.0, 1.25, 2.2, 2.15, "系统设计", ["完成分层架构、业务模型、数据库与模块划分", "明确图片知识库、文档知识库、检索与问答接口"])
    add_card(slide, 5.35, 1.25, 2.2, 2.15, "原型开发", ["完成 FastAPI + Vue3 原型", "实现图片/文档入库、检索、RAG问答、会话管理"])
    add_card(slide, 7.7, 1.25, 1.3, 2.15, "实验准备", ["评测脚本", "对比方法", "结果分析"])
    add_textbox(slide, 0.8, 4.0, 8.2, 0.28, "当前阶段性结论", font_size=18, bold=True, color=COLOR_NAVY)
    add_bullets(
        slide,
        0.8,
        4.35,
        8.1,
        2.0,
        [
            "系统已形成“数据接入 -> 语义化 -> 检索 -> RAG问答 -> 前端展示”的闭环原型。",
            "实验框架、图表与结果分析材料已具备，可支撑中期汇报与后续论文实验章节。",
            "文档解析稳定性、检索效率和微调效果仍需在下一阶段继续打磨。",
        ],
        font_size=16,
    )
    add_footer(slide, "已完成工作以仓库当前代码、docs 文档和实验数据为依据")


def build_architecture_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统架构与代码落地", "Architecture")
    layers = [
        ("展示层", ["Vue3 前端", "KnowledgeBase / DocumentKB / Search / Chat"]),
        ("应用层", ["知识管理服务", "ChatSession 持久化", "RAG 编排与流式接口"]),
        ("检索层", ["向量检索", "BM25 / 混合检索", "重排序与上下文压缩"]),
        ("语义化层", ["图像描述生成", "文档解析/切块", "查询改写与意图识别"]),
        ("存储层", ["SQLite 元数据", "ChromaDB 向量库", "本地文件存储"]),
    ]
    top = 1.1
    for idx, (title, bullets) in enumerate(layers):
        add_card(slide, 0.85, top + idx * 1.08, 3.0, 0.88, title, bullets)
    add_textbox(slide, 4.35, 1.2, 4.8, 0.28, "与代码目录对应", font_size=18, bold=True, color=COLOR_NAVY)
    add_bullets(
        slide,
        4.35,
        1.55,
        4.5,
        3.8,
        [
            "后端入口：backend/main.py -> app.main.create_app()",
            "路由模块：api/routers/kb.py、docs.py、search.py、chat.py",
            "业务编排：application/knowledge_management.py、chat_service.py",
            "RAG 适配层：langchain_integration/adapters.py",
            "评测模块：evaluation/run_offline_eval.py、run_unidoc_eval.py 等",
            "前端页面：frontend/src/views/*.vue 与 api/*.ts",
        ],
        font_size=15,
    )
    add_textbox(slide, 4.35, 5.85, 4.5, 0.85, "当前代码层面已经明确支持：图片知识库、文档知识库、文本/图像检索、会话化RAG问答，以及离线评测脚本。", font_size=15)
    add_footer(slide, "架构说明基于 backend/app、frontend/src 和 evaluation 目录的真实实现")


def build_current_status_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "当前系统实现情况", "Current Status")
    add_card(slide, 0.7, 1.25, 4.1, 2.2, "后端能力", [
        "图片上传后自动生成描述，并同步写入 SQLite、Chroma 和文件存储。",
        "支持 PDF 上传、后台解析、进度查询、结果查看、重处理和删除。",
        "支持文本到图像、图像到图像检索，以及带会话历史的 RAG 问答。",
    ])
    add_card(slide, 5.0, 1.25, 4.1, 2.2, "前端能力", [
        "已实现图片知识库、文档知识库、检索页面和聊天页面。",
        "支持筛选、编辑、预览、重处理、引用来源查看和会话切换。",
        "前后端接口已联调，具备完整演示路径。",
    ])
    add_card(slide, 0.7, 3.8, 4.1, 2.1, "检索增强链路", [
        "已集成查询改写、多查询扩展、混合检索、重排序和上下文压缩。",
        "支持基于意图识别的轻量化路由，增强问答过程可追溯性。",
    ])
    add_card(slide, 5.0, 3.8, 4.1, 2.1, "实验与材料", [
        "已沉淀检索/生成评测脚本、分析图表和微调训练日志。",
        "中期报告、系统设计报告和技术调研文档已基本成型。",
    ])
    add_footer(slide, "系统已达到“可运行原型 + 可展示 + 可评测”的中期阶段目标")


def build_ui_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统界面展示", "UI")
    images = [
        ("图片知识库", ASSET_DIR / "image1.png"),
        ("文档知识库", ASSET_DIR / "image2.png"),
        ("图文检索服务", ASSET_DIR / "image3.png"),
        ("RAG 问答服务", ASSET_DIR / "image4.png"),
    ]
    slots = [(0.55, 1.05), (5.0, 1.05), (0.55, 4.0), (5.0, 4.0)]
    for (caption, path), (left, top) in zip(images, slots):
        frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(top), Inches(4.0), Inches(2.55))
        frame.fill.solid()
        frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
        frame.line.color.rgb = COLOR_BORDER
        add_picture_contain(slide, path, left + 0.08, top + 0.08, 3.84, 2.15)
        add_caption(slide, left, top + 2.28, 4.0, caption)
    add_footer(slide, "四个核心页面分别对应知识管理、检索和多轮问答场景")


def build_feature_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "核心功能与优化点", "Features")
    add_card(slide, 0.7, 1.2, 2.7, 1.45, "图片语义化", ["MLLM 生成结构化描述", "构建可检索图像知识库"])
    add_card(slide, 3.6, 1.2, 2.7, 1.45, "文档解析", ["PDF 后台解析与切块", "页面图像与文本片段统一管理"])
    add_card(slide, 6.5, 1.2, 2.7, 1.45, "双路检索", ["文本检图", "以图搜图统一走文本检索链路"])
    add_card(slide, 0.7, 3.0, 2.7, 1.55, "会话化RAG", ["会话创建/重命名/删除", "消息历史与来源快照持久化"])
    add_card(slide, 3.6, 3.0, 2.7, 1.55, "检索增强", ["Query Rewrite / Multi-Query", "Hybrid Retrieval / Rerank / Context Compression"])
    add_card(slide, 6.5, 3.0, 2.7, 1.55, "可解释性", ["引用来源展示", "检索轨迹与过程信息返回"])
    add_textbox(slide, 0.8, 5.15, 8.6, 1.0, "当前版本的重点不是单点功能演示，而是把图片、文档、检索、RAG问答、会话管理和评测链路连成统一系统。", font_size=16)
    add_footer(slide, "优化项已在 backend/app/langchain_integration 与 evaluation 中形成对应实现和实验材料")


def build_retrieval_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "检索实验结果（不含 OCR 基线）", "Retrieval Evaluation")
    charts = [
        ("CRM 领域指标", RETRIEVAL_FIG_DIR / "retrieval_crm_without_ocr_v2.png"),
        ("Energy 领域指标", RETRIEVAL_FIG_DIR / "retrieval_energy_without_ocr_v2.png"),
        ("两种方法延迟对比", RETRIEVAL_FIG_DIR / "retrieval_latency_without_ocr_v2.png"),
        ("分领域指标热力图", RETRIEVAL_FIG_DIR / "retrieval_heatmap_without_ocr_v2.png"),
    ]
    slots = [(0.45, 1.0), (5.0, 1.0), (0.45, 4.0), (5.0, 4.0)]
    for (caption, path), (left, top) in zip(charts, slots):
        add_picture_contain(slide, path, left, top, 4.1, 2.35)
        add_caption(slide, left, top + 2.27, 4.1, caption)
    add_textbox(slide, 0.55, 6.45, 8.8, 0.65, "结论：Proposed 在 CRM 和 Energy 场景中接近甚至超过 baseline_clip；在 Finance 与 Cross-domain 场景下仍弱于 baseline_clip，说明复杂表格页和跨领域混合检索仍是后续重点。", font_size=14)
    add_footer(slide, "代表性指标：CRM Recall@10 0.9768 vs 0.9472；Energy mAP@10 0.9589 vs 0.9367；延迟整体在约 500-580ms 区间")


def build_generation_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "生成实验结果（不含 OCR 基线）", "Generation Evaluation")
    charts = [
        ("总体指标总览", GEN_FIG_DIR / "generation_overall_metrics_without_ocr_v2.png"),
        ("分领域答案正确性", GEN_FIG_DIR / "generation_domain_correctness_without_ocr_v2.png"),
        ("本文方法相对基线提升", GEN_FIG_DIR / "generation_improvement_without_ocr_v2.png"),
        ("相关性与忠实度对比", GEN_FIG_DIR / "generation_relevancy_faithfulness_without_ocr_v2.png"),
    ]
    slots = [(0.45, 1.0), (5.0, 1.0), (0.45, 4.0), (5.0, 4.0)]
    for (caption, path), (left, top) in zip(charts, slots):
        add_picture_contain(slide, path, left, top, 4.1, 2.35)
        add_caption(slide, left, top + 2.27, 4.1, caption)
    add_textbox(slide, 0.55, 6.42, 8.8, 0.68, "总体来看，proposed 的答案正确性为 0.86，与 baseline_clip 的 0.87 基本接近，但显著优于 no_rag 的 0.24；同时图像上下文相关性达到 0.9286，说明检索到的视觉证据与问题更匹配。", font_size=14)
    add_footer(slide, "当前主文展示重点为 proposed 与 baseline_clip 的对比，已按你的要求忽略 OCR 异常基线")


def build_finetune_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "模型微调探索", "Fine-tuning")
    add_textbox(slide, 0.7, 1.0, 3.9, 0.28, "训练设置", font_size=18, bold=True, color=COLOR_NAVY)
    add_bullets(
        slide,
        0.7,
        1.35,
        3.8,
        3.2,
        [
            "基于 UniDoc-Bench 图像与描述文本构建监督微调数据集。",
            "训练样本 888 条，验证样本 47 条。",
            "基于 Qwen3.5-4B 进行 LoRA 微调，共训练 140 step，耗时约 5 小时 20 分。",
            "step-75 的 eval_loss 最低，为 0.5892；最终 train_loss 约 0.5194，末轮 eval_loss 约 0.5915。",
        ],
        font_size=15,
    )
    add_textbox(slide, 0.7, 4.85, 3.9, 0.28, "阶段性观察", font_size=18, bold=True, color=COLOR_NAVY)
    add_bullets(
        slide,
        0.7,
        5.2,
        3.8,
        1.4,
        [
            "部分复杂页面的格式遵循、关键词补充和局部细节描述能力有所提升。",
            "但从阶段性人工评估看，整体效果尚未稳定超过 base 模型，仍需继续优化数据与训练策略。",
        ],
        font_size=14,
    )
    add_picture_contain(slide, FINETUNE_FIG_DIR / "finetune_learning_curves_v3.png", 4.7, 1.15, 4.2, 5.4)
    add_caption(slide, 4.75, 6.55, 4.1, "微调学习曲线")
    add_footer(slide, "微调部分当前更适合作为“可行性验证 + 后续优化方向”，而不是最终定稿结论")


def build_problem_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "当前存在问题", "Open Issues")
    add_card(slide, 0.75, 1.25, 4.0, 1.7, "1. 文档解析稳定性仍需增强", ["复杂 PDF、表格页面、跨页内容和长文档处理仍存在波动。"])
    add_card(slide, 5.0, 1.25, 4.0, 1.7, "2. 检索与问答效率有优化空间", ["多步检索增强与外部模型调用会拉长整体响应时间。"])
    add_card(slide, 0.75, 3.25, 4.0, 1.7, "3. 强基线场景下仍有差距", ["Finance / Cross-domain 场景中 baseline_clip 仍然更强，说明复杂版面语义建模还需提升。"])
    add_card(slide, 5.0, 3.25, 4.0, 1.7, "4. 微调效果尚未收敛", ["step-75 的损失更优，但整体人工评估尚未稳定优于 base。"])
    add_textbox(slide, 0.85, 5.55, 8.1, 0.7, "说明：OCR 基线当前存在异常，因此本次汇报不将其作为正式对比对象，后续会单独修复索引与评测流程。", font_size=15, color=COLOR_RED)
    add_footer(slide, "中期阶段重点是识别系统短板，并为后续实验与工程优化留出空间")


def build_plan_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "下一阶段工作计划", "Next Steps")
    steps = [
        ("01", "修复 OCR 基线并补跑实验", "重新构建异常索引，恢复对照实验完整性。"),
        ("02", "继续优化图像描述与微调策略", "围绕数据质量、LoRA 参数和输出格式稳定性持续迭代。"),
        ("03", "完善文档知识库与多模态 RAG 链路", "提升表格、长文档和复杂页面场景的稳定性与可解释性。"),
        ("04", "加强测试与演示版本打磨", "补足边界处理、异常提示与前端交互细节。"),
        ("05", "推进论文撰写与材料整理", "同步沉淀系统设计、实验结果、图表和答辩素材。"),
    ]
    top = 1.15
    for idx, (no, title, desc) in enumerate(steps):
        left = 0.7 if idx % 2 == 0 else 5.1
        row = idx // 2
        card_top = top + row * 1.85
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(card_top), Inches(3.8), Inches(1.45))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_LIGHT
        shape.line.color.rgb = COLOR_BORDER
        add_textbox(slide, left + 0.15, card_top + 0.12, 0.4, 0.25, no, font_size=18, bold=True, color=COLOR_RED)
        add_textbox(slide, left + 0.55, card_top + 0.12, 3.0, 0.25, title, font_size=15, bold=True, color=COLOR_NAVY)
        add_textbox(slide, left + 0.55, card_top + 0.46, 3.0, 0.6, desc, font_size=13)
    add_footer(slide, "目标是将当前原型进一步推进为“可验证、可分析、可答辩展示”的毕业设计成果")


def build_end_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_textbox(slide, 1.0, 2.15, 7.5, 0.8, "汇报完毕  谢谢老师", font_size=30, bold=True, color=COLOR_NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 2.2, 3.35, 5.2, 0.5, "张琪  |  东北大学计算机2202", font_size=18, color=COLOR_MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, 2.8, 4.05, 4.2, 0.4, "2026年4月", font_size=16, color=COLOR_RED, align=PP_ALIGN.CENTER)
    add_footer(slide, "中期汇报")


def main() -> None:
    prs = Presentation(str(TEMPLATE_PATH))
    remove_all_slides(prs)

    build_title_slide(prs)
    build_outline_slide(prs)
    build_background_slide(prs)
    build_progress_slide(prs)
    build_architecture_slide(prs)
    build_current_status_slide(prs)
    build_ui_slide(prs)
    build_feature_slide(prs)
    build_retrieval_slide(prs)
    build_generation_slide(prs)
    build_finetune_slide(prs)
    build_problem_slide(prs)
    build_plan_slide(prs)
    build_end_slide(prs)

    prs.save(str(OUTPUT_PATH))
    print(f"saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
