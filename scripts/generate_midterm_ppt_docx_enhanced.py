from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
ASSET_DIR = DOCS_DIR / "midterm_ppt_assets" / "docx_media"
DOCX_PATH = DOCS_DIR / "张琪-中期报告-v2.docx"
TEMPLATE_PATH = DOCS_DIR / "PPT模版.pptx"
OUTPUT_PATH = DOCS_DIR / "张琪-中期汇报-展示增强版.pptx"

COLOR_NAVY = RGBColor(16, 49, 97)
COLOR_RED = RGBColor(181, 45, 54)
COLOR_TEXT = RGBColor(38, 38, 38)
COLOR_MUTED = RGBColor(92, 101, 116)
COLOR_LIGHT = RGBColor(245, 247, 251)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BORDER = RGBColor(214, 222, 234)
COLOR_TINT = RGBColor(232, 238, 248)
COLOR_SOFT_RED = RGBColor(252, 239, 240)


def ensure_docx_media() -> list[Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    with ZipFile(DOCX_PATH) as archive:
        media_names = sorted(
            [name for name in archive.namelist() if name.startswith("word/media/")],
            key=lambda item: int(Path(item).stem.replace("image", "")),
        )
        paths: list[Path] = []
        for name in media_names:
            target = ASSET_DIR / Path(name).name
            if not target.exists():
                target.write_bytes(archive.read(name))
            paths.append(target)
    return paths


def remove_all_slides(prs: Presentation) -> None:
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def add_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[0])


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    font_size: int = 18,
    bold: bool = False,
    color: RGBColor = COLOR_TEXT,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bullets(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    items: list[str],
    font_size: int = 16,
    color: RGBColor = COLOR_TEXT,
    spacing: int = 8,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.bullet = True


def add_top_banner(slide, title: str, subtitle: str | None = None) -> None:
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        0,
        0,
        Inches(10),
        Inches(0.72),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_NAVY
    bar.line.fill.background()
    add_textbox(slide, 0.55, 0.14, 7.5, 0.34, title, font_size=24, bold=True, color=COLOR_WHITE)
    if subtitle:
        add_textbox(slide, 8.0, 0.16, 1.5, 0.25, subtitle, font_size=10, color=RGBColor(225, 232, 242), align=PP_ALIGN.RIGHT)


def add_footer(slide, text: str) -> None:
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.45),
        Inches(7.15),
        Inches(9.1),
        Inches(0.02),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_RED
    line.line.fill.background()
    add_textbox(slide, 0.55, 7.18, 6.0, 0.18, text, font_size=9, color=COLOR_MUTED)


def add_card(slide, left: float, top: float, width: float, height: float, title: str, bullets: list[str], tint: RGBColor = COLOR_LIGHT) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = tint
    shape.line.color.rgb = COLOR_BORDER
    add_textbox(slide, left + 0.16, top + 0.12, width - 0.3, 0.25, title, font_size=16, bold=True, color=COLOR_NAVY)
    add_bullets(slide, left + 0.12, top + 0.42, width - 0.24, height - 0.48, bullets, font_size=13, spacing=4)


def add_stat_chip(slide, left: float, top: float, width: float, height: float, big: str, label: str) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_TINT
    shape.line.color.rgb = COLOR_BORDER
    add_textbox(slide, left, top + 0.12, width, 0.32, big, font_size=22, bold=True, color=COLOR_RED, align=PP_ALIGN.CENTER)
    add_textbox(slide, left + 0.08, top + 0.52, width - 0.16, 0.26, label, font_size=10, color=COLOR_MUTED, align=PP_ALIGN.CENTER)


def contain_size(path: Path, max_w: float, max_h: float) -> tuple[float, float]:
    with Image.open(path) as img:
        width, height = img.size
    ratio = min(max_w / width, max_h / height)
    return width * ratio, height * ratio


def add_picture_contain(slide, path: Path, left: float, top: float, width: float, height: float) -> None:
    img_w, img_h = contain_size(path, Inches(width), Inches(height))
    x = Inches(left) + (Inches(width) - img_w) / 2
    y = Inches(top) + (Inches(height) - img_h) / 2
    slide.shapes.add_picture(str(path), x, y, width=img_w, height=img_h)


def add_image_frame(slide, left: float, top: float, width: float, height: float, tint: RGBColor = COLOR_WHITE) -> None:
    frame = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = tint
    frame.line.color.rgb = COLOR_BORDER


def add_caption(slide, left: float, top: float, width: float, text: str) -> None:
    add_textbox(slide, left, top, width, 0.24, text, font_size=10, color=COLOR_MUTED, align=PP_ALIGN.CENTER)


def build_title_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_LIGHT
    bg.line.fill.background()

    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(6.8), Inches(0), Inches(3.2), Inches(7.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_TINT
    accent.line.fill.background()

    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.05), Inches(0.16), Inches(1.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_RED
    bar.line.fill.background()

    add_textbox(slide, 1.0, 1.0, 6.0, 0.35, "东北大学毕业设计（论文）中期汇报", font_size=16, color=COLOR_RED)
    add_textbox(slide, 1.0, 1.55, 5.9, 1.25, "基于语义描述桥接的\n多模态RAG系统设计与实现", font_size=28, bold=True, color=COLOR_NAVY)
    add_textbox(slide, 1.05, 3.15, 4.8, 0.45, "Academic Enhanced Edition", font_size=18, color=COLOR_MUTED)
    add_textbox(slide, 1.05, 4.35, 3.9, 0.95, "汇报人：张琪\n学号：20225928\n班级：计算机2202", font_size=17)
    add_stat_chip(slide, 7.15, 1.45, 2.05, 1.15, "14", "页展示增强版")
    add_stat_chip(slide, 7.15, 2.95, 2.05, 1.15, "DOCX", "唯一内容来源")
    add_stat_chip(slide, 7.15, 4.45, 2.05, 1.15, "2026.04", "中期答辩")
    add_footer(slide, "基于《张琪-中期报告-v2.docx》内容重新组织的学术增强版")


def build_outline_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "汇报内容", "Outline")
    add_textbox(slide, 0.7, 1.0, 2.0, 0.25, "内容脉络", font_size=16, bold=True, color=COLOR_RED)
    timeline = [
        ("01", "前期工作", "文献学习与系统开发"),
        ("02", "系统展示", "界面与核心功能"),
        ("03", "实验评估", "检索与生成结果"),
        ("04", "模型微调", "训练设置与阶段性观察"),
        ("05", "问题与计划", "存在问题和下一阶段安排"),
    ]
    for idx, (num, title, desc) in enumerate(timeline):
        top = 1.45 + idx * 1.05
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.95), Inches(top + 0.03), Inches(0.3), Inches(0.3))
        dot.fill.solid()
        dot.fill.fore_color.rgb = COLOR_RED if idx == 0 else COLOR_NAVY
        dot.line.fill.background()
        if idx < len(timeline) - 1:
            line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1.08), Inches(top + 0.35), Inches(0.04), Inches(0.75))
            line.fill.solid()
            line.fill.fore_color.rgb = COLOR_BORDER
            line.line.fill.background()
        add_textbox(slide, 1.45, top, 0.45, 0.22, num, font_size=13, bold=True, color=COLOR_RED)
        add_textbox(slide, 1.92, top - 0.02, 2.2, 0.24, title, font_size=18, bold=True, color=COLOR_NAVY)
        add_textbox(slide, 1.95, top + 0.3, 2.8, 0.25, desc, font_size=11, color=COLOR_MUTED)

    add_card(slide, 5.25, 1.35, 3.9, 1.55, "展示增强目标", ["强化页面层次与讲述节奏", "提升图表页结论先读性", "保留正式答辩气质"], tint=COLOR_TINT)
    add_card(slide, 5.25, 3.1, 3.9, 1.55, "内容约束", ["仅使用 docx 原文和内嵌图表", "不混用实验数据目录版本", "不改动已有两版成品"], tint=COLOR_LIGHT)
    add_textbox(slide, 5.35, 5.2, 3.7, 0.7, "这版重点不是增加内容，而是把已有材料组织成更适合现场答辩的阅读顺序。", font_size=15)
    add_footer(slide, "目录页改为“主线 + 目标 + 约束”的学术汇报结构")


def build_literature_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "前期文献阅读与资料学习", "Literature")
    add_textbox(slide, 0.7, 0.98, 3.2, 0.3, "阶段结论", font_size=16, bold=True, color=COLOR_RED)
    add_textbox(slide, 0.7, 1.35, 4.2, 0.9, "已完成多模态检索、多模态RAG、向量数据库与评测方法的系统调研，为原型设计和实验验证提供理论基础。", font_size=19, bold=True, color=COLOR_NAVY)
    add_card(slide, 0.7, 2.55, 2.6, 2.45, "重点学习内容", ["CLIP、BLIP-2、LLaVA", "Qwen-VL、MuRAG", "多模态RAG综述与基准"], tint=COLOR_LIGHT)
    add_card(slide, 3.5, 2.55, 2.6, 2.45, "方法与工具", ["RAGAs 等自动评测方法", "向量数据库与嵌入模型", "检索相关性与忠实性分析"], tint=COLOR_TINT)
    add_card(slide, 6.3, 2.55, 2.6, 2.45, "形成认识", ["检索相关性", "答案忠实性", "上下文利用", "向量存储与召回机制"], tint=COLOR_LIGHT)
    add_textbox(slide, 0.82, 5.45, 8.1, 0.72, "相较普通 bullet 罗列，这一页突出“学了什么、用来做什么、形成了什么认识”三层逻辑，更适合中期答辩表述。", font_size=15)
    add_footer(slide, "内容依据：中期报告“1.前期文献阅读与资料学习”")


def build_design_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统设计与开发", "Design")
    add_textbox(slide, 0.7, 1.0, 2.6, 0.28, "阶段结论", font_size=16, bold=True, color=COLOR_RED)
    add_textbox(slide, 0.7, 1.35, 4.1, 0.75, "已完成多模态 RAG 系统的整体设计与原型开发，形成可运行的前后端闭环框架。", font_size=20, bold=True, color=COLOR_NAVY)
    add_stat_chip(slide, 6.9, 1.05, 1.9, 1.05, "FastAPI", "后端框架")
    add_stat_chip(slide, 6.9, 2.35, 1.9, 1.05, "Vue3", "前端框架")
    add_stat_chip(slide, 6.9, 3.65, 1.9, 1.05, "Vite", "工程工具")
    add_card(slide, 0.72, 2.55, 2.55, 2.0, "设计工作", ["业务模型", "功能模型", "数据模型", "数据库与模块设计"], tint=COLOR_LIGHT)
    add_card(slide, 3.48, 2.55, 2.8, 2.0, "已完成功能", ["图片知识库", "文档知识库", "图文检索服务", "RAG 问答服务", "基础管理功能"], tint=COLOR_TINT)
    add_textbox(slide, 0.78, 5.15, 5.55, 0.85, "这一阶段的关键不是单个模块开发，而是把设计、功能和技术选型收束成一个能够实际运行与展示的系统原型。", font_size=15)
    add_footer(slide, "内容依据：中期报告“2.系统设计与开发”")


def build_ui_knowledge_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统界面展示（一）", "UI")
    add_textbox(slide, 0.7, 0.98, 2.8, 0.28, "知识管理入口", font_size=16, bold=True, color=COLOR_RED)
    add_image_frame(slide, 0.7, 1.4, 4.2, 4.8, tint=COLOR_WHITE)
    add_picture_contain(slide, media[0], 0.8, 1.5, 4.0, 4.2)
    add_caption(slide, 0.75, 5.8, 4.1, "图表 1 图片知识库界面")
    add_image_frame(slide, 5.1, 1.4, 4.2, 4.8, tint=COLOR_WHITE)
    add_picture_contain(slide, media[1], 5.2, 1.5, 4.0, 4.2)
    add_caption(slide, 5.15, 5.8, 4.1, "图表 2 文档知识库界面")
    add_card(slide, 6.15, 0.96, 3.0, 0.95, "页面价值", ["完成图片与文档两类知识的接入与组织"], tint=COLOR_TINT)
    add_footer(slide, "图片与标题均来自中期报告内嵌图表 1-2")


def build_ui_service_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "系统界面展示（二）", "UI")
    add_textbox(slide, 0.7, 0.98, 3.0, 0.28, "应用服务入口", font_size=16, bold=True, color=COLOR_RED)
    add_image_frame(slide, 0.7, 1.4, 4.2, 4.8, tint=COLOR_WHITE)
    add_picture_contain(slide, media[2], 0.8, 1.5, 4.0, 4.2)
    add_caption(slide, 0.75, 5.8, 4.1, "图表 3 图文检索服务界面")
    add_image_frame(slide, 5.1, 1.4, 4.2, 4.8, tint=COLOR_WHITE)
    add_picture_contain(slide, media[3], 5.2, 1.5, 4.0, 4.2)
    add_caption(slide, 5.15, 5.8, 4.1, "图表 4 RAG 问答服务")
    add_card(slide, 6.12, 0.96, 3.05, 0.95, "页面价值", ["完成从检索到多模态问答的业务闭环"], tint=COLOR_TINT)
    add_footer(slide, "图片与标题均来自中期报告内嵌图表 3-4")


def build_features_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "核心功能实现与优化", "Functions")
    add_textbox(slide, 0.7, 1.0, 4.4, 0.72, "系统已基本完成开题报告中的主要功能目标，并在检索增强链路上继续做了完整化集成。", font_size=20, bold=True, color=COLOR_NAVY)
    add_card(slide, 0.7, 2.0, 2.7, 1.75, "基础能力", ["图像描述生成", "结构化信息入库", "向量化存储"], tint=COLOR_LIGHT)
    add_card(slide, 3.6, 2.0, 2.7, 1.75, "检索能力", ["文本检索图片", "以图搜图", "双路检索机制"], tint=COLOR_TINT)
    add_card(slide, 6.5, 2.0, 2.7, 1.75, "问答能力", ["RAG 回答内容", "引用来源", "检索过程信息"], tint=COLOR_LIGHT)
    add_card(slide, 0.7, 4.0, 2.7, 1.75, "增强机制", ["查询改写", "多查询扩展", "混合检索"], tint=COLOR_TINT)
    add_card(slide, 3.6, 4.0, 2.7, 1.75, "结果优化", ["重排序", "上下文压缩", "轻量化路由"], tint=COLOR_LIGHT)
    add_card(slide, 6.5, 4.0, 2.7, 1.75, "阶段价值", ["结果可追溯", "过程可查看", "能按意图切换模式"], tint=COLOR_TINT)
    add_footer(slide, "内容依据：中期报告“3.核心功能实现与优化”")


def build_eval_intro_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "实验评估总述", "Evaluation")
    add_textbox(slide, 0.7, 1.02, 3.2, 0.25, "阶段结论", font_size=16, bold=True, color=COLOR_RED)
    add_textbox(slide, 0.7, 1.35, 4.4, 0.85, "已完成检索器评估和生成器评估两类实验，阶段性结果初步验证了所提方法的有效性和竞争力。", font_size=20, bold=True, color=COLOR_NAVY)
    add_stat_chip(slide, 5.5, 1.2, 1.6, 1.05, "2类", "已完成实验")
    add_stat_chip(slide, 7.3, 1.2, 1.6, 1.05, "proposed", "主方法")
    add_card(slide, 0.7, 2.65, 4.0, 2.45, "检索实验结论", ["在 CRM 和 energy 领域表现较好", "在 finance 和 cross-domain 场景下仍有优化空间"], tint=COLOR_LIGHT)
    add_card(slide, 4.95, 2.65, 4.0, 2.45, "生成实验结论", ["回答正确性已接近 baseline_clip", "图像上下文相关性达到最高"], tint=COLOR_TINT)
    add_textbox(slide, 0.8, 5.55, 8.05, 0.65, "这一页先给答辩老师一个全局认识：后面两页图表不是零散结果，而是围绕“检索有效、生成可用、但仍有弱项”这一主线展开。", font_size=15)
    add_footer(slide, "内容依据：中期报告“4.实验评估”")


def build_retrieval_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "检索实验结果", "Retrieval")
    add_card(slide, 0.65, 0.95, 8.7, 0.85, "页内结论", ["proposed 在多个领域接近甚至超过 baseline_clip；优势主要体现在 CRM 和 energy，弱项集中在 finance 与 cross-domain。"], tint=COLOR_TINT)
    items = [
        ("CRM领域相关检索指标", media[4]),
        ("Energy领域相关检索指标", media[5]),
        ("两种方法检索延迟对比", media[6]),
        ("分领域指标热力图", media[7]),
    ]
    slots = [(0.45, 1.95), (5.0, 1.95), (0.45, 4.55), (5.0, 4.55)]
    for (caption, path), (left, top) in zip(items, slots):
        add_image_frame(slide, left, top, 4.1, 2.18)
        add_picture_contain(slide, path, left + 0.06, top + 0.06, 3.98, 1.85)
        add_caption(slide, left, top + 1.93, 4.1, caption)
    add_footer(slide, "内容依据：中期报告图表 5-8 与对应文字结论")


def build_generation_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "生成实验结果", "Generation")
    add_card(slide, 0.65, 0.95, 8.7, 0.85, "页内结论", ["proposed 的回答正确性已接近最佳 baseline_clip，同时在图像上下文相关性上达到最高。"], tint=COLOR_TINT)
    add_stat_chip(slide, 7.65, 5.55, 1.2, 0.95, "0.86", "回答正确性")
    add_stat_chip(slide, 6.2, 5.55, 1.2, 0.95, "0.9286", "图像上下文相关性")
    placements = [
        ("图表 9 三种方法生成器评测总览", media[8], 0.42, 1.95, 2.95, 1.95),
        ("图表 10 三种方法分领域答案正确性", media[9], 3.55, 1.95, 2.95, 1.95),
        ("图表 11 本文方法相对基线的提升", media[10], 6.68, 1.95, 2.9, 1.95),
        ("图表 12 分问题类型答案正确性", media[11], 1.18, 4.15, 3.1, 1.95),
        ("图表 13 相关性与忠实度对比", media[12], 4.95, 4.15, 3.1, 1.95),
    ]
    for caption, path, left, top, width, height in placements:
        add_image_frame(slide, left, top, width, height + 0.18)
        add_picture_contain(slide, path, left + 0.05, top + 0.05, width - 0.1, height - 0.15)
        add_caption(slide, left, top + height - 0.02, width, caption)
    add_footer(slide, "内容依据：中期报告图表 9-13 与对应文字结论")


def build_finetune_slide(prs: Presentation, media: list[Path]) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "模型微调", "Fine-tuning")
    add_textbox(slide, 0.7, 1.0, 3.7, 0.72, "已完成监督微调数据构建、训练执行和阶段性评估，当前结果证明了专用图像描述微调的可行性。", font_size=19, bold=True, color=COLOR_NAVY)
    add_stat_chip(slide, 0.72, 2.05, 1.2, 1.0, "888", "训练样本")
    add_stat_chip(slide, 2.05, 2.05, 1.2, 1.0, "47", "验证样本")
    add_stat_chip(slide, 3.38, 2.05, 1.2, 1.0, "140", "训练 step")
    add_stat_chip(slide, 4.71, 2.05, 1.45, 1.0, "0.5892", "最低验证损失")
    add_card(slide, 0.7, 3.45, 3.7, 2.15, "阶段性观察", ["部分复杂页面上提升了格式遵循、关键词补充和局部细节描述能力。", "但仍存在输出冗长、重复等问题，整体效果还需继续优化。"], tint=COLOR_LIGHT)
    add_image_frame(slide, 4.85, 1.65, 4.1, 4.8)
    add_picture_contain(slide, media[13], 4.95, 1.75, 3.9, 4.15)
    add_caption(slide, 5.0, 6.0, 3.8, "图表 14 微调学习曲线")
    add_footer(slide, "内容依据：中期报告“5.模型微调”")


def build_issues_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "存在问题", "Issues")
    add_textbox(slide, 0.7, 1.0, 3.0, 0.25, "当前短板", font_size=16, bold=True, color=COLOR_RED)
    add_card(slide, 0.7, 1.35, 4.0, 1.45, "1. OCR 基线检索异常", ["OCR 基线索引构建不完整，导致 Recall@K、MRR 等指标失真。"], tint=COLOR_SOFT_RED)
    add_card(slide, 5.0, 1.35, 4.0, 1.45, "2. 模型微调效果不够理想", ["当前微调虽已完成阶段性训练，但整体效果提升仍不够稳定。"], tint=COLOR_LIGHT)
    add_card(slide, 0.7, 3.15, 4.0, 1.45, "3. 系统响应速度仍有待提升", ["串行检索优化策略与 IO 操作叠加，导致问答延迟较大。"], tint=COLOR_LIGHT)
    add_card(slide, 5.0, 3.15, 4.0, 1.45, "4. 工程化与鲁棒性仍存在不足", ["异常处理、恢复机制和工程化规范方面仍有提升空间。"], tint=COLOR_SOFT_RED)
    add_textbox(slide, 0.8, 5.35, 8.1, 0.65, "这一页强调的不是问题越多越好，而是阶段性地识别出真正影响后续实验与系统完善的关键阻碍。", font_size=15)
    add_footer(slide, "内容依据：中期报告“存在问题”")


def build_plan_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    add_top_banner(slide, "下一阶段工作计划", "Plan")
    add_textbox(slide, 0.7, 1.0, 3.2, 0.25, "问题对应措施", font_size=16, bold=True, color=COLOR_RED)
    rows = [
        ("01", "修复 OCR 基线并重跑实验", "恢复对照实验完整性"),
        ("02", "优化模型微调与图像描述能力", "提升格式遵循和语义概括效果"),
        ("03", "完善多模态 RAG 链路并提升性能", "兼顾检索精度与问答效率"),
        ("04", "加强测试并完善系统体验", "补足异常处理和前端交互细节"),
        ("05", "推进论文初稿撰写与材料整理", "为后续定稿和答辩做准备"),
    ]
    for idx, (num, title, desc) in enumerate(rows):
        top = 1.45 + idx * 1.02
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(top), Inches(8.2), Inches(0.74))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_LIGHT if idx % 2 == 0 else COLOR_TINT
        shape.line.color.rgb = COLOR_BORDER
        add_textbox(slide, 1.0, top + 0.12, 0.45, 0.2, num, font_size=16, bold=True, color=COLOR_RED)
        add_textbox(slide, 1.55, top + 0.1, 4.65, 0.22, title, font_size=16, bold=True, color=COLOR_NAVY)
        add_textbox(slide, 6.4, top + 0.12, 2.0, 0.18, desc, font_size=11, color=COLOR_MUTED, align=PP_ALIGN.RIGHT)
    add_footer(slide, "内容依据：中期报告“下一阶段工作计划”")


def build_end_slide(prs: Presentation) -> None:
    slide = add_slide(prs)
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_LIGHT
    bg.line.fill.background()
    accent = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(6.75), Inches(10), Inches(0.75))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_NAVY
    accent.line.fill.background()
    add_textbox(slide, 1.15, 2.1, 7.6, 0.8, "汇报完毕  谢谢老师", font_size=30, bold=True, color=COLOR_NAVY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 2.3, 3.3, 5.4, 0.4, "张琪  |  东北大学计算机2202", font_size=18, color=COLOR_MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, 2.9, 4.02, 4.1, 0.35, "2026年4月", font_size=16, color=COLOR_RED, align=PP_ALIGN.CENTER)
    add_footer(slide, "学术增强版中期汇报")


def main() -> None:
    media = ensure_docx_media()
    if len(media) != 14:
        raise ValueError(f"expected 14 docx media images, got {len(media)}")

    prs = Presentation(str(TEMPLATE_PATH))
    remove_all_slides(prs)

    build_title_slide(prs)
    build_outline_slide(prs)
    build_literature_slide(prs)
    build_design_slide(prs)
    build_ui_knowledge_slide(prs, media)
    build_ui_service_slide(prs, media)
    build_features_slide(prs)
    build_eval_intro_slide(prs)
    build_retrieval_slide(prs, media)
    build_generation_slide(prs, media)
    build_finetune_slide(prs, media)
    build_issues_slide(prs)
    build_plan_slide(prs)
    build_end_slide(prs)

    prs.save(str(OUTPUT_PATH))
    print(f"saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
