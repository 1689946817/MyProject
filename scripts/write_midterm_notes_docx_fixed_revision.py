from __future__ import annotations

from pathlib import Path

from pptx import Presentation


PPT_PATH = Path(r"E:\BiShe\Code2\MyProject\docs\张琪-中期汇报-docx修正版.pptx")

NOTES = [
    """各位老师好，我是张琪。我的毕业设计题目是“基于语义描述桥接的多模态RAG系统设计与实现”。下面我将结合当前系统实现和阶段性实验结果，从前期工作、系统展示、实验评估、模型微调、存在问题和下一阶段计划几个方面做中期汇报。""",
    """本次汇报大致分成六个部分。前面先说明前期调研和系统开发进展，中间展示系统界面与核心能力，后面集中汇报检索实验、生成实验和模型微调结果，最后再归纳当前存在的问题以及下一阶段的工作安排。""",
    """前一阶段，我围绕课题做了比较系统的文献阅读和资料学习。重点关注的方向包括多模态检索、检索增强生成、多模态大模型、向量数据库以及 RAG 自动评测方法。通过这些调研，我对系统的可选技术路线、评价指标和关键问题形成了比较明确的认识，也为后续原型开发和实验设计打下了基础。""",
    """从系统建设进度来看，前一阶段已经完成了整体设计和原型开发工作，包括业务模型、功能模型和数据模型的建立，以及数据库和模块层面的设计。技术选型上，后端采用 FastAPI，前端采用 Vue3、TypeScript 和 Vite。目前图片知识库、文档知识库、图文检索服务、RAG 问答服务以及基础管理功能都已经完成，系统具备了较完整的运行框架。""",
    """先看图片知识库部分。这个模块主要承担图片数据的接入、管理和语义化组织，是后续文本检图和以图搜图能力的基础。通过这个界面可以看到，目前系统已经具备了实际的数据管理入口，而不只是停留在概念设计层面。""",
    """接着看文档知识库。这个模块面向文档上传、解析和知识管理，承担的是文档侧的知识组织能力。图片知识库和文档知识库共同构成了系统多模态知识接入的底座，也为后续检索和问答服务提供了统一的数据支撑。""",
    """再往后是图文检索服务。这里对应两类核心能力，一类是文本检索图片，另一类是以图搜图。也就是说，系统已经能够把不同形式的用户输入统一纳入检索链路，并把检索结果以可视化方式输出。""",
    """在应用层面，RAG 问答服务已经可以基于检索到的多模态证据生成回答，并返回引用来源和相关过程信息。到这里为止，知识接入、检索和问答三个关键环节已经串联起来，形成了一个可运行的多模态 RAG 原型。""",
    """从核心能力来看，当前系统已经实现图像描述生成、结构化入库和向量化存储，也实现了文本检图和以图搜图的双路检索机制。同时，系统支持基于 RAG 的多模态问答，并加入了查询改写、多查询扩展、混合检索、重排序、上下文压缩等优化环节，使整体链路更加完整。""",
    """除了基础检索和问答能力，我还补充实现了意图识别模块。这个模块把当前支持的意图类型做了归纳，核心作用是根据用户问题自动切换执行模式，比如进入检索流程、问答流程或者其他对应链路。它的价值在于让系统不再只是被动执行单一流程，而是具备一定的任务路由和自适应能力。""",
    """在系统功能打通之后，前一阶段还完成了两类实验评估，分别是检索器评估和生成器评估。这样做的目的，是让系统不仅能运行起来，还能通过实验结果来验证方法有效性、分析当前优势，并定位后续仍需优化的部分。""",
    """先看检索实验中表现较好的两个领域，也就是 CRM 和 Energy。CRM 场景里，proposed 在 Recall@10 上优于 baseline，大约是 0.977 对 0.947，MRR 和 mAP@10 也基本接近或略有优势；Energy 场景中，proposed 的 Recall@1、Recall@10、MRR 和 mAP@10 都略高于 baseline。说明当前方法在这两个领域已经体现出比较稳定的竞争力。""",
    """接下来补充 Education、Finance、Healthcare 和 Legal 四个领域。Education 基本上是两种方法持平；Healthcare 上 proposed 在 Recall@10 略有优势，但在 MRR 和 mAP@10 上稍弱；Finance 和 Legal 场景中，baseline_clip 整体仍然更强。这个结果说明 proposed 并不是在所有领域都占优，它的表现仍然受到页面类型和领域特征的影响。""",
    """Commerce & Manu. 和 Construction 这两个领域的结果进一步说明了这一点。在这两个场景里，baseline_clip 的整体表现依然更好，尤其在 Recall@1、MRR 和 mAP@10 上更明显。也就是说，当前 proposed 方法在一些复杂工业或结构化页面场景下，还需要继续优化。""",
    """跨领域场景的难度会更高一些。从图里可以看到，proposed 在四个指标上都略低于 baseline_clip，尤其是 mAP@10 和 MRR 的差距更明显。这比较符合预期，因为跨领域混合检索对语义泛化能力要求更高，也是后续需要重点突破的方向之一。""",
    """把前面几页分领域结果综合起来看，按领域汇总的 Recall@1 和 MRR 能更清楚地反映整体趋势。proposed 在 CRM 和 Energy 上相对突出，在部分领域和 baseline 接近，但在 Commerce、Construction、Finance 这类场景下仍有差距。这一页实际上是把前面的分散结果收束成一个整体判断。""",
    """除了看指标本身，我还分析了首个命中结果的排名分布。可以看到，baseline_clip 有更高比例的结果直接命中在第 1 名，而 proposed 更多落在前 2 到 5 名区间。换句话说，当前方法虽然能把正确结果稳定召回到较靠前的位置，但在把最优结果直接顶到第一名这件事上，排序能力还有继续优化的空间。""",
    """再看检索延迟。两种方法的整体延迟都在五百多毫秒量级，proposed 的平均延迟略低于 baseline，大约是 562 毫秒对 580 毫秒；中位延迟两者非常接近，proposed 还略高一点。说明当前方法在效果提升的同时，并没有带来明显不可接受的时延代价。""",
    """最后用热力图把各个领域、各项指标的结果放在一起比较。热力图很直观地说明，proposed 的亮点主要集中在 CRM 和 Energy，而 baseline 在 Finance、Construction、Cross-domain 等场景仍然更稳。到这里，检索实验的结论就比较完整了：当前方法已经有明显亮点，但还不是全面领先。""",
    """下面看生成实验的总体结果。proposed 的回答正确性达到 0.86，与 baseline_clip 的 0.87 已经非常接近，而 no_rag 只有 0.24，差距非常明显。同时，在图像上下文相关性上，proposed 达到 0.929，是三种方法里最高的。这个结果说明，引入检索增强后，回答质量有明显提升，而且 proposed 在视觉证据匹配上具有一定优势。""",
    """按领域来看，proposed 在 Commerce、Education、Finance 和 Legal 这些领域可以达到与 baseline_clip 持平，甚至在 Education 上更优；在 CRM 和 Construction 上略弱于 baseline；在 Energy 和 Healthcare 上两者比较接近。整体来看，proposed 的生成效果在多数领域已经具备较强竞争力，但仍然存在领域间表现不均衡的问题。""",
    """如果把 proposed 分别和两个基线做差值比较，结果会更清楚。相对于 no_rag，proposed 在回答正确性、图像上下文相关性和图像忠实度上的提升都很明显，尤其图像上下文相关性提升接近 0.93。相对于 baseline_clip，proposed 在图像上下文相关性和答案相关性上略有增益，但在回答正确性和图像忠实度上还没有明显超过它。""",
    """进一步按问题类型来看，在 Casual 和 Temporal 这两类问题上，proposed 与 baseline_clip 基本持平；在 Factual 类型上，proposed 还略优；但在 Comparison 和 Summarization 这类更复杂的问题上，proposed 仍然稍弱。这说明当前方法在直接事实型问题上的表现已经比较稳定，但在更依赖综合归纳和比较能力的题型上还需要继续提升。""",
    """再看相关性和忠实度的对比。Answer Relevancy 上三种方法都很高，差距不大；Image Context Relevancy 上，proposed 略高于 baseline_clip；而 Image Faithfulness 上，baseline_clip 仍然稍强一些。也就是说，proposed 更擅长检索到和问题更匹配的视觉上下文，但在最终生成内容对图像证据的忠实表达上，还存在继续优化的空间。""",
    """模型微调部分也已经完成了阶段性探索。目前基于 UniDoc-Bench 图像及描述文本构建了监督微调数据集，其中训练样本 888 条、验证样本 47 条，并完成了基于 Qwen3.5-4B 的 LoRA 微调训练，共 140 个 step。训练损失整体持续下降，验证损失在 step-75 左右达到最低的 0.5892，说明训练过程是有效的；但后期验证损失改善有限，也表明当前微调效果还没有完全稳定。""",
    """综合来看，目前存在四个比较突出的短板。第一，OCR 基线索引构建不完整，导致相关指标失真；第二，模型微调效果还不够稳定，与预期目标仍有差距；第三，系统在响应速度方面还有优化空间；第四，在异常处理、恢复机制和工程化规范方面仍需继续完善。这些问题基本对应了当前系统从实验到工程落地两条线上的主要不足。""",
    """下一阶段的工作计划主要围绕五个方向展开。首先是修复 OCR 基线问题并重跑相关实验，保证实验对照完整；其次是继续优化模型微调和图像描述能力；第三是完善多模态 RAG 链路并降低系统延迟；第四是加强测试和系统体验优化；最后是同步推进论文初稿撰写和材料整理，为后续论文定稿和答辩做好准备。""",
    """我的汇报到这里结束。整体来看，前一阶段我已经完成了文献调研、系统原型开发、检索与生成实验以及模型微调探索等工作，系统已经具备了较完整的运行框架。接下来我会继续围绕实验完善、性能优化和论文写作三个方向推进。请各位老师批评指正。""",
]


def write_notes() -> None:
    prs = Presentation(str(PPT_PATH))
    if len(prs.slides) != len(NOTES):
        raise ValueError(f"slide count {len(prs.slides)} does not match notes count {len(NOTES)}")

    for slide, note_text in zip(prs.slides, NOTES):
        notes_shape = None
        for shape in slide.notes_slide.shapes:
            if getattr(shape, "is_placeholder", False) and shape.placeholder_format.idx == 3:
                notes_shape = shape
                break
        if notes_shape is None:
            raise RuntimeError("notes placeholder not found")

        text_frame = notes_shape.text_frame
        text_frame.clear()
        text_frame.text = note_text

    prs.save(str(PPT_PATH))
    print(f"notes written: {PPT_PATH}")


if __name__ == "__main__":
    write_notes()
